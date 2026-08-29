"""
Sinh file Excel template để nhập hàng loạt: Trạm + QR Alias + Thông số.

Cách dùng:
    cd backend
    python tools/make_import_template.py
    # -> tạo ../import-template/cau_hinh_template.xlsx

Điền xong file đó rồi import bằng:
    python tools/import_config.py ../import-template/cau_hinh_template.xlsx
"""
import os
import sys

try:  # in tiếng Việt được trên console Windows (cp1252)
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# backend/ lên sys.path để có thể import config tĩnh làm ví dụ
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from services.stations_config import STATIONS, QR_ALIAS_MAP, STATION_PARAMS  # noqa: E402

OUT_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "import-template")
OUT_PATH = os.path.join(OUT_DIR, "cau_hinh_template.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
EXAMPLE_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _add_rows(ws, rows, fill=None):
    for r in rows:
        ws.append(r)
        if fill:
            for c in range(1, len(r) + 1):
                ws.cell(row=ws.max_row, column=c).fill = fill
                ws.cell(row=ws.max_row, column=c).border = BORDER


def build():
    wb = Workbook()

    # ----- Sheet Hướng dẫn -----
    g = wb.active
    g.title = "Hướng dẫn"
    instructions = [
        ["TEMPLATE NHẬP HÀNG LOẠT — QR Checklist", ""],
        ["", ""],
        ["Điền dữ liệu vào 3 sheet bên dưới, KHÔNG đổi tên cột (hàng 1) và KHÔNG đổi tên sheet.", ""],
        ["Hàng tô màu xanh là VÍ DỤ — xóa đi rồi điền dữ liệu thật, hoặc sửa đè lên.", ""],
        ["Import bằng lệnh: python tools/import_config.py <đường-dẫn-file>.xlsx", ""],
        ["", ""],
        ["SHEET 'Trạm'", "tọa độ + bán kính geofence + checklist trạm thuộc về"],
        ["  name", "Tên trạm (vd TK-5201A). Bắt buộc, sẽ tự viết HOA."],
        ["  lat / lng", "Tọa độ. Bắt buộc. Lấy từ Google Maps (nhấn giữ → copy)."],
        ["  radius", "Bán kính cho phép (mét). Để trống = 50."],
        ["  checklist_types", "Các checklist, cách nhau dấu phẩy (vd: tank,routine). Để trống = chưa gán."],
        ["  active", "TRUE/FALSE. Để trống = TRUE (đang dùng)."],
        ["", ""],
        ["SHEET 'QR Alias'", "map nội dung QR thực tế dán tại trạm → tên trạm"],
        ["  qr_content", "Nội dung QR (vd 052-LI-022B). Bắt buộc, duy nhất."],
        ["  station_name", "Tên trạm tương ứng (phải có trong sheet Trạm). Bắt buộc."],
        ["  note", "Ghi chú tùy chọn."],
        ["", ""],
        ["SHEET 'Thông số'", "các thông số vận hành cần nhập sau khi quét — MỖI trạm NHIỀU dòng"],
        ["  station_name", "Tên trạm. Bắt buộc."],
        ["  tag", "Mã thiết bị (vd 052-PG-038). Tùy chọn nhưng nên có để đối chiếu."],
        ["  param_label", "Tên thông số (vd Áp suất, Nhiệt độ). Bắt buộc."],
        ["  param_unit", "Đơn vị (mm, bar, °C, A...). Đơn vị 'Yes/No' → nhập Y/N."],
        ["  param_low", "Ngưỡng dưới (số). Để trống nếu không cảnh báo cận dưới."],
        ["  param_high", "Ngưỡng trên (số). Để trống nếu không cảnh báo cận trên."],
        ["  sort_order", "Thứ tự hiển thị (số nhỏ lên trước). Để trống = 0."],
        ["  active", "TRUE/FALSE. Để trống = TRUE. Đặt FALSE để ẩn thông số cũ."],
        ["", ""],
        ["KHÓA TRÙNG khi import lại (upsert — không tạo bản sao):", ""],
        ["  Trạm:     theo 'name'", ""],
        ["  QR Alias: theo 'qr_content'", ""],
        ["  Thông số: theo (station_name + tag + param_label)", ""],
        ["Import KHÔNG xóa dữ liệu cũ. Để gỡ 1 thông số: đặt active=FALSE rồi import lại.", ""],
    ]
    for row in instructions:
        g.append(row)
    g["A1"].font = Font(bold=True, size=14, color="1F4E78")
    for r in range(7, g.max_row + 1):
        a = g.cell(row=r, column=1)
        if a.value and a.value.startswith("SHEET"):
            a.font = Font(bold=True, color="1F4E78")
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 90

    # ----- Sheet Trạm -----
    st = wb.create_sheet("Trạm")
    st.append(["name", "lat", "lng", "radius", "checklist_types", "active"])
    _style_header(st, 6)
    examples = []
    for name, cfg in list(STATIONS.items())[:3]:
        examples.append([name, cfg.get("lat"), cfg.get("lng"), cfg.get("radius", 50), "tank", "TRUE"])
    _add_rows(st, examples, EXAMPLE_FILL)
    _autosize(st, [16, 14, 14, 10, 22, 10])

    # ----- Sheet QR Alias -----
    al = wb.create_sheet("QR Alias")
    al.append(["qr_content", "station_name", "note"])
    _style_header(al, 3)
    ex_al = [[k, v, ""] for k, v in list(QR_ALIAS_MAP.items())[:3]]
    _add_rows(al, ex_al, EXAMPLE_FILL)
    _autosize(al, [22, 18, 30])

    # ----- Sheet Thông số -----
    pa = wb.create_sheet("Thông số")
    pa.append(["station_name", "tag", "param_label", "param_unit",
               "param_low", "param_high", "sort_order", "active"])
    _style_header(pa, 8)
    ex_pa = [
        ["TK-5211A", "052-LI-042B", "Tank level", "mm", 100, 4000, 1, "TRUE"],
        ["PUMP_STATION_6", "052-PG-038", "Áp suất", "bar", 2, 8, 1, "TRUE"],
        ["PUMP_STATION_6", "052-PG-038", "Nhiệt độ", "°C", "", 75, 2, "TRUE"],
        ["PUMP_STATION_6", "052-PG-038", "Chạy/Dừng", "Yes/No", "", "", 3, "TRUE"],
    ]
    _add_rows(pa, ex_pa, EXAMPLE_FILL)
    _autosize(pa, [18, 16, 18, 12, 12, 12, 11, 10])

    os.makedirs(os.path.abspath(OUT_DIR), exist_ok=True)
    wb.save(os.path.abspath(OUT_PATH))
    print(f"[OK] Da tao template: {os.path.abspath(OUT_PATH)}")


if __name__ == "__main__":
    build()
