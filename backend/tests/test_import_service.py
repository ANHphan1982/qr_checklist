"""
TDD — services/import_service.py: import hàng loạt Trạm/Alias/Thông số từ
workbook Excel (tách từ tools/import_config.py để route admin dùng chung).

Yêu cầu:
- import_workbook(session, wb) đọc 3 sheet 'Trạm' / 'QR Alias' / 'Thông số',
  upsert như tool CLI cũ, trả về dict kết quả có cấu trúc (không print):
    {"stations"|"aliases"|"params": {"created","updated","skipped"} | None,
     "warnings": [...], "missing_sheets": [...]}
- Sheet thiếu → giá trị None + tên sheet trong missing_sheets
- Hàng thiếu trường bắt buộc → skipped + 1 dòng warning
- build_template_workbook() trả về Workbook template 4 sheet (Hướng dẫn,
  Trạm, QR Alias, Thông số) để route cho tải về
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from unittest.mock import MagicMock
from openpyxl import Workbook

from models import Station, QrAlias, StationParam


STATION_HEADERS = ["name", "lat", "lng", "radius", "checklist_types", "active"]
ALIAS_HEADERS = ["qr_content", "station_name", "note"]
PARAM_HEADERS = ["station_name", "tag", "param_label", "param_unit",
                 "param_low", "param_high", "sort_order", "active"]


def _wb(stations=None, aliases=None, params=None):
    """Workbook in-memory với các sheet được truyền (None = không tạo sheet)."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in [("Trạm", STATION_HEADERS, stations),
                                 ("QR Alias", ALIAS_HEADERS, aliases),
                                 ("Thông số", PARAM_HEADERS, params)]:
        if rows is None:
            continue
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append(r)
    return wb


def _session():
    """Session giả dispatch theo model; mặc định DB rỗng (first() -> None)."""
    sess = MagicMock()
    chains = {}

    def q(model):
        if model not in chains:
            m = MagicMock()
            m.filter.return_value.first.return_value = None
            m.filter.return_value.filter.return_value.first.return_value = None
            chains[model] = m
        return chains[model]

    sess.query.side_effect = q
    return sess, chains


class TestImportStations:
    def test_creates_new_station_normalized_upper(self):
        from services.import_service import import_workbook
        sess, _ = _session()
        result = import_workbook(sess, _wb(stations=[["tk-1", 10.1, 20.2, 100, "tank", "TRUE"]]))
        assert result["stations"] == {"created": 1, "updated": 0, "skipped": 0}
        added = sess.add.call_args[0][0]
        assert isinstance(added, Station)
        assert added.name == "TK-1"
        assert added.radius == 100
        assert added.checklist_types == ["tank"]
        assert added.checklist_type == "tank"

    def test_updates_existing_station(self):
        from services.import_service import import_workbook
        sess, chains = _session()
        st = MagicMock()
        m = MagicMock()
        m.filter.return_value.first.return_value = st
        chains[Station] = m
        result = import_workbook(sess, _wb(stations=[["TK-1", 11.0, 21.0, 80, "", ""]]))
        assert result["stations"] == {"created": 0, "updated": 1, "skipped": 0}
        assert st.lat == 11.0
        assert st.lng == 21.0
        assert st.radius == 80
        sess.add.assert_not_called()

    def test_skips_station_missing_coords_with_warning(self):
        from services.import_service import import_workbook
        sess, _ = _session()
        result = import_workbook(sess, _wb(stations=[["TK-1", None, 20.2, "", "", ""]]))
        assert result["stations"] == {"created": 0, "updated": 0, "skipped": 1}
        assert len(result["warnings"]) == 1
        assert "TK-1" in result["warnings"][0]
        sess.add.assert_not_called()


class TestImportAliases:
    def test_creates_new_alias(self):
        from services.import_service import import_workbook
        sess, _ = _session()
        result = import_workbook(sess, _wb(aliases=[["052-LI-042B", "tk-1", "ghi chú"]]))
        assert result["aliases"] == {"created": 1, "updated": 0, "skipped": 0}
        added = sess.add.call_args[0][0]
        assert isinstance(added, QrAlias)
        assert added.qr_content == "052-LI-042B"
        assert added.station_name == "TK-1"

    def test_skips_alias_missing_station(self):
        from services.import_service import import_workbook
        sess, _ = _session()
        result = import_workbook(sess, _wb(aliases=[["052-LI-042B", "", ""]]))
        assert result["aliases"] == {"created": 0, "updated": 0, "skipped": 1}
        assert len(result["warnings"]) == 1


class TestImportParams:
    def test_creates_new_param_with_defaults(self):
        from services.import_service import import_workbook
        sess, _ = _session()
        result = import_workbook(sess, _wb(params=[["tk-1", "", "Tank level", "", "", "", "", ""]]))
        assert result["params"] == {"created": 1, "updated": 0, "skipped": 0}
        added = sess.add.call_args[0][0]
        assert isinstance(added, StationParam)
        assert added.station_name == "TK-1"
        assert added.param_label == "Tank level"
        assert added.param_unit == "mm"
        assert added.sort_order == 0
        assert added.active is True

    def test_updates_existing_param(self):
        from services.import_service import import_workbook
        sess, chains = _session()
        sp = MagicMock()
        m = MagicMock()
        m.filter.return_value.filter.return_value.first.return_value = sp
        chains[StationParam] = m
        result = import_workbook(
            sess, _wb(params=[["TK-1", "052-LI-042B", "Tank level", "mm", 100, 4000, 2, "TRUE"]]))
        assert result["params"] == {"created": 0, "updated": 1, "skipped": 0}
        assert sp.param_low == 100.0
        assert sp.param_high == 4000.0
        assert sp.sort_order == 2
        sess.add.assert_not_called()

    def test_skips_param_missing_label(self):
        from services.import_service import import_workbook
        sess, _ = _session()
        result = import_workbook(sess, _wb(params=[["TK-1", "", "", "mm", "", "", "", ""]]))
        assert result["params"] == {"created": 0, "updated": 0, "skipped": 1}
        assert len(result["warnings"]) == 1


class TestSheetHandling:
    def test_missing_sheets_reported(self):
        from services.import_service import import_workbook
        sess, _ = _session()
        result = import_workbook(sess, _wb(stations=[["TK-1", 1.0, 2.0, "", "", ""]]))
        assert result["stations"] is not None
        assert result["aliases"] is None
        assert result["params"] is None
        assert set(result["missing_sheets"]) == {"QR Alias", "Thông số"}

    def test_all_sheets_missing(self):
        from services.import_service import import_workbook
        sess, _ = _session()
        wb = Workbook()  # chỉ có sheet mặc định, không sheet nào đúng tên
        result = import_workbook(sess, wb)
        assert result["stations"] is None
        assert result["aliases"] is None
        assert result["params"] is None
        assert set(result["missing_sheets"]) == {"Trạm", "QR Alias", "Thông số"}

    def test_no_warnings_when_all_rows_valid(self):
        from services.import_service import import_workbook
        sess, _ = _session()
        result = import_workbook(sess, _wb(stations=[["TK-1", 1.0, 2.0, "", "", ""]]))
        assert result["warnings"] == []


class TestTemplateBuilder:
    def test_builds_template_with_4_sheets(self):
        from services.import_service import build_template_workbook
        wb = build_template_workbook()
        assert wb.sheetnames == ["Hướng dẫn", "Trạm", "QR Alias", "Thông số"]

    def test_template_station_headers(self):
        from services.import_service import build_template_workbook
        wb = build_template_workbook()
        headers = [c.value for c in wb["Trạm"][1]]
        assert headers == STATION_HEADERS
