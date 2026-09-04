"""
TDD — Admin API import cấu hình từ Excel:

- POST /api/admin/import-config: body JSON {file_base64, dry_run}.
  Auth X-Admin-Key; DB thiếu → 503; thiếu file → 400; file hỏng → 400;
  file không có sheet nào đúng tên → 400.
  dry_run=true → chạy import rồi ROLLBACK (xem trước); ngược lại COMMIT.
  Response = kết quả import_workbook + cờ dry_run.
- GET /api/admin/import-template: trả file template .xlsx (auth như trên).
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import base64
from io import BytesIO

import pytest
from unittest.mock import patch, MagicMock
from openpyxl import Workbook

import config as cfg


def _xlsx_b64(sheets=("Trạm",)):
    """File xlsx in-memory: mỗi sheet 1 hàng dữ liệu hợp lệ, trả về base64."""
    wb = Workbook()
    wb.remove(wb.active)
    if "Trạm" in sheets:
        ws = wb.create_sheet("Trạm")
        ws.append(["name", "lat", "lng", "radius", "checklist_types", "active"])
        ws.append(["TK-1", 10.1, 20.2, 100, "tank", "TRUE"])
    if "QR Alias" in sheets:
        ws = wb.create_sheet("QR Alias")
        ws.append(["qr_content", "station_name", "note"])
        ws.append(["052-LI-042B", "TK-1", ""])
    if not wb.sheetnames:
        wb.create_sheet("Khac")
    buf = BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode("ascii")


def _session():
    sess = MagicMock()
    sess.__enter__ = MagicMock(return_value=sess)
    sess.__exit__ = MagicMock(return_value=False)
    chain = MagicMock()
    chain.filter.return_value.first.return_value = None
    chain.filter.return_value.filter.return_value.first.return_value = None
    sess.query.return_value = chain
    return sess


@pytest.fixture
def client():
    with patch.object(cfg, "SessionLocal", None):
        from app import app as flask_app
        flask_app.config["TESTING"] = True
        yield flask_app.test_client()


def _post(client, sess, body, key="test-secret"):
    with patch("routes.admin.ADMIN_SECRET", "test-secret"), \
         patch("routes.admin.SessionLocal", (lambda: sess) if sess else None):
        return client.post("/api/admin/import-config", json=body,
                           headers={"X-Admin-Key": key})


class TestImportConfigRoute:
    def test_requires_admin_key(self, client):
        resp = _post(client, _session(), {"file_base64": _xlsx_b64()}, key="wrong")
        assert resp.status_code == 401

    def test_503_when_db_unavailable(self, client):
        resp = _post(client, None, {"file_base64": _xlsx_b64()})
        assert resp.status_code == 503

    def test_400_when_missing_file(self, client):
        resp = _post(client, _session(), {})
        assert resp.status_code == 400

    def test_400_when_file_not_xlsx(self, client):
        b64 = base64.b64encode(b"khong phai xlsx").decode("ascii")
        resp = _post(client, _session(), {"file_base64": b64})
        assert resp.status_code == 400

    def test_400_when_no_known_sheet(self, client):
        resp = _post(client, _session(), {"file_base64": _xlsx_b64(sheets=())})
        assert resp.status_code == 400

    def test_dry_run_rolls_back_and_returns_preview(self, client):
        sess = _session()
        resp = _post(client, sess, {"file_base64": _xlsx_b64(sheets=("Trạm", "QR Alias")),
                                    "dry_run": True})
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["dry_run"] is True
        assert data["stations"] == {"created": 1, "updated": 0, "skipped": 0}
        assert data["aliases"] == {"created": 1, "updated": 0, "skipped": 0}
        assert data["params"] is None
        sess.rollback.assert_called_once()
        sess.commit.assert_not_called()

    def test_real_import_commits(self, client):
        sess = _session()
        resp = _post(client, sess, {"file_base64": _xlsx_b64()})
        assert resp.status_code == 200
        data = resp.get_json()
        assert data["dry_run"] is False
        assert data["stations"] == {"created": 1, "updated": 0, "skipped": 0}
        sess.commit.assert_called_once()
        sess.rollback.assert_not_called()


class TestImportTemplateRoute:
    def test_requires_admin_key(self, client):
        with patch("routes.admin.ADMIN_SECRET", "test-secret"):
            resp = client.get("/api/admin/import-template",
                              headers={"X-Admin-Key": "wrong"})
        assert resp.status_code == 401

    def test_downloads_xlsx_template(self, client):
        with patch("routes.admin.ADMIN_SECRET", "test-secret"):
            resp = client.get("/api/admin/import-template",
                              headers={"X-Admin-Key": "test-secret"})
        assert resp.status_code == 200
        assert resp.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        # nội dung phải là workbook đọc được, đủ 4 sheet
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(resp.data))
        assert wb.sheetnames == ["Hướng dẫn", "Trạm", "QR Alias", "Thông số"]
