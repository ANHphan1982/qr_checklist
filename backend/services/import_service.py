"""
Import hàng loạt Trạm + QR Alias + Thông số từ workbook Excel template.

Dùng chung cho:
- CLI: tools/import_config.py (giữ cách dùng cũ)
- API: POST /api/admin/import-config (upload từ trang Admin)

Upsert (không tạo bản sao khi import lại):
    Trạm     -> theo name
    QR Alias -> theo qr_content
    Thông số -> theo (station_name + tag + param_label)

Import KHÔNG xóa dữ liệu cũ. Để gỡ 1 thông số: đặt active=FALSE rồi import lại.
Caller tự quyết định commit (import thật) hay rollback (dry-run).
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import Station, QrAlias, StationParam, resolve_checklist_list

SHEET_STATIONS = "Trạm"
SHEET_ALIASES = "QR Alias"
SHEET_PARAMS = "Thông số"


# --------------------------------------------------------------------------- #
# Helpers parse ô Excel
# --------------------------------------------------------------------------- #
def _s(v):
    """Chuỗi đã strip; None/ô trống -> ''."""
    if v is None:
        return ""
    return str(v).strip()


def _float_or_none(v):
    s = _s(v)
    if s == "":
        return None
    try:
        return float(s.replace(",", "."))
    except ValueError:
        return None


def _int_or(v, default=0):
    s = _s(v)
    if s == "":
        return default
    try:
        return int(float(s))
    except ValueError:
        return default


def _bool_or(v, default=True):
    s = _s(v).lower()
    if s == "":
        return default
    return s not in {"0", "false", "no", "n", "khong", "không", "tat", "tắt", "off", "x"}


def _rows(ws):
    """Yield dict {header: value} cho từng hàng dữ liệu (bỏ hàng tiêu đề + hàng rỗng)."""
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return
    headers = [_s(h) for h in rows[0]]
    for raw in rows[1:]:
        if raw is None or all(c is None or _s(c) == "" for c in raw):
            continue
        yield {headers[i]: (raw[i] if i < len(raw) else None) for i in range(len(headers))}


# --------------------------------------------------------------------------- #
# Import từng sheet — trả {"created","updated","skipped"}, cảnh báo dồn vào warnings
# --------------------------------------------------------------------------- #
def import_stations(session, ws, warnings):
    created = updated = skipped = 0
    for row in _rows(ws):
        name = _s(row.get("name")).upper()
        lat = _float_or_none(row.get("lat"))
        lng = _float_or_none(row.get("lng"))
        if not name or lat is None or lng is None:
            warnings.append(f"[Trạm] Bỏ qua hàng thiếu name/lat/lng: {row}")
            skipped += 1
            continue
        radius = _int_or(row.get("radius"), 50)
        active = _bool_or(row.get("active"), True)
        ct_raw = [c.strip() for c in _s(row.get("checklist_types")).split(",") if c.strip()]
        checklist_types = resolve_checklist_list(ct_raw, None)

        st = session.query(Station).filter(Station.name == name).first()
        if st:
            st.lat, st.lng, st.radius, st.active = lat, lng, radius, active
            if ct_raw:  # chỉ ghi đè checklist khi ô có giá trị (tránh xóa nhầm)
                st.checklist_types = checklist_types
                st.checklist_type = checklist_types[0] if checklist_types else None
            updated += 1
        else:
            st = Station(name=name, lat=lat, lng=lng, radius=radius, active=active,
                         checklist_types=checklist_types,
                         checklist_type=checklist_types[0] if checklist_types else None)
            session.add(st)
            created += 1
    return {"created": created, "updated": updated, "skipped": skipped}


def import_aliases(session, ws, warnings):
    created = updated = skipped = 0
    for row in _rows(ws):
        qr = _s(row.get("qr_content"))
        station = _s(row.get("station_name")).upper()
        note = _s(row.get("note")) or None
        if not qr or not station:
            warnings.append(f"[QR Alias] Bỏ qua hàng thiếu qr_content/station_name: {row}")
            skipped += 1
            continue
        al = session.query(QrAlias).filter(QrAlias.qr_content == qr).first()
        if al:
            al.station_name, al.note = station, note
            updated += 1
        else:
            session.add(QrAlias(qr_content=qr, station_name=station, note=note))
            created += 1
    return {"created": created, "updated": updated, "skipped": skipped}


def import_params(session, ws, warnings):
    created = updated = skipped = 0
    for row in _rows(ws):
        station = _s(row.get("station_name")).upper()
        label = _s(row.get("param_label"))
        if not station or not label:
            warnings.append(f"[Thông số] Bỏ qua hàng thiếu station_name/param_label: {row}")
            skipped += 1
            continue
        tag = _s(row.get("tag")) or None
        unit = _s(row.get("param_unit")) or "mm"
        low = _float_or_none(row.get("param_low"))
        high = _float_or_none(row.get("param_high"))
        sort_order = _int_or(row.get("sort_order"), 0)
        active = _bool_or(row.get("active"), True)

        q = session.query(StationParam).filter(
            StationParam.station_name == station,
            StationParam.param_label == label,
        )
        # khóa trùng gồm cả tag (NULL khớp NULL)
        q = q.filter(StationParam.tag == tag) if tag else q.filter(StationParam.tag.is_(None))
        sp = q.first()
        if sp:
            sp.param_unit, sp.param_low, sp.param_high = unit, low, high
            sp.sort_order, sp.active = sort_order, active
            updated += 1
        else:
            session.add(StationParam(station_name=station, tag=tag, param_label=label,
                                     param_unit=unit, param_low=low, param_high=high,
                                     sort_order=sort_order, active=active))
            created += 1
    return {"created": created, "updated": updated, "skipped": skipped}


_SHEET_HANDLERS = [
    (SHEET_STATIONS, "stations", import_stations),
    (SHEET_ALIASES, "aliases", import_aliases),
    (SHEET_PARAMS, "params", import_params),
]


def import_workbook(session, wb):
    """Chạy import cả 3 sheet trên session đang mở. KHÔNG commit/rollback —
    caller quyết định (dry-run rollback, import thật commit)."""
    result = {"warnings": [], "missing_sheets": []}
    for sheet, key, handler in _SHEET_HANDLERS:
        if sheet in wb.sheetnames:
            result[key] = handler(session, wb[sheet], result["warnings"])
        else:
            result[key] = None
            result["missing_sheets"].append(sheet)
    return result


# --------------------------------------------------------------------------- #
# Template builder (chuyển từ tools/make_import_template.py)
# --------------------------------------------------------------------------- #
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_EXAMPLE_FILL = PatternFill("solid", fgColor="E2EFDA")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _autosize(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _add_rows(ws, rows, fill=None):
    for r in rows:
        ws.append(r)
        if fill:
            for c in range(1, len(r) + 1):
                ws.cell(row=ws.max_row, column=c).fill = fill
                ws.cell(row=ws.max_row, column=c).border = _BORDER


def build_template_workbook():
    """Workbook template nhập hàng loạt: Hướng dẫn + Trạm + QR Alias + Thông số."""
    from services.stations_config import STATIONS, QR_ALIAS_MAP

    wb = Workbook()

    # ----- Sheet Hướng dẫn -----
    g = wb.active
    g.title = "Hướng dẫn"
    instructions = [
        ["TEMPLATE NHẬP HÀNG LOẠT — QR Checklist", ""],
        ["", ""],
        ["Điền dữ liệu vào 3 sheet bên dưới, KHÔNG đổi tên cột (hàng 1) và KHÔNG đổi tên sheet.", ""],
        ["Hàng tô màu xanh là VÍ DỤ — xóa đi rồi điền dữ liệu thật, hoặc sửa đè lên.", ""],
        ["Import: trang Admin → tab Import, hoặc lệnh python tools/import_config.py <file>.xlsx", ""],
        ["", ""],
        ["SHEET 'Trạm'", "tọa độ + bán kính geofence + checklist trạm thuộc về"],
        ["  name", "Tên trạm (vd TK-5201A). Bắt buộc, sẽ tự viết HOA."],
        ["  lat / lng", "Tọa độ. Bắt buộc. Lấy từ Google Maps (nhấn giữ → copy)."],
        ["  radius", "Bán kính cho phép (mét). Để trống = 50."],
        ["  checklist_types", "Các checklist, cách nhau dấu phẩy (vd: tank,routine). Để trống = chưa gán."],
        ["  active", "TRUE/FALSE. Để trống = TRUE (đang dùng)."],
        ["", ""],
        ["SHEET 'QR Alias'", "map nội dung QR thực tế dán tại trạm → tên trạm"],
        ["  qr_content", "Nội dung QR (vd 052-LI-022B). Bắt buộc, duy nhất."],
        ["  station_name", "Tên trạm tương ứng (phải có trong sheet Trạm). Bắt buộc."],
        ["  note", "Ghi chú tùy chọn."],
        ["", ""],
        ["SHEET 'Thông số'", "các thông số vận hành cần nhập sau khi quét — MỖI trạm NHIỀU dòng"],
        ["  station_name", "Tên trạm. Bắt buộc."],
        ["  tag", "Mã thiết bị (vd 052-PG-038). Tùy chọn nhưng nên có để đối chiếu."],
        ["  param_label", "Tên thông số (vd Áp suất, Nhiệt độ). Bắt buộc."],
        ["  param_unit", "Đơn vị (mm, bar, °C, A...). Đơn vị 'Yes/No' → nhập Y/N."],
        ["  param_low", "Ngưỡng dưới (số). Để trống nếu không cảnh báo cận dưới."],
        ["  param_high", "Ngưỡng trên (số). Để trống nếu không cảnh báo cận trên."],
        ["  sort_order", "Thứ tự hiển thị (số nhỏ lên trước). Để trống = 0."],
        ["  active", "TRUE/FALSE. Để trống = TRUE. Đặt FALSE để ẩn thông số cũ."],
        ["", ""],
        ["KHÓA TRÙNG khi import lại (upsert — không tạo bản sao):", ""],
        ["  Trạm:     theo 'name'", ""],
        ["  QR Alias: theo 'qr_content'", ""],
        ["  Thông số: theo (station_name + tag + param_label)", ""],
        ["Import KHÔNG xóa dữ liệu cũ. Để gỡ 1 thông số: đặt active=FALSE rồi import lại.", ""],
    ]
    for row in instructions:
        g.append(row)
    g["A1"].font = Font(bold=True, size=14, color="1F4E78")
    for r in range(7, g.max_row + 1):
        a = g.cell(row=r, column=1)
        if a.value and a.value.startswith("SHEET"):
            a.font = Font(bold=True, color="1F4E78")
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 90

    # ----- Sheet Trạm -----
    st = wb.create_sheet(SHEET_STATIONS)
    st.append(["name", "lat", "lng", "radius", "checklist_types", "active"])
    _style_header(st, 6)
    examples = []
    for name, cfg in list(STATIONS.items())[:3]:
        examples.append([name, cfg.get("lat"), cfg.get("lng"), cfg.get("radius", 50), "tank", "TRUE"])
    _add_rows(st, examples, _EXAMPLE_FILL)
    _autosize(st, [16, 14, 14, 10, 22, 10])

    # ----- Sheet QR Alias -----
    al = wb.create_sheet(SHEET_ALIASES)
    al.append(["qr_content", "station_name", "note"])
    _style_header(al, 3)
    ex_al = [[k, v, ""] for k, v in list(QR_ALIAS_MAP.items())[:3]]
    _add_rows(al, ex_al, _EXAMPLE_FILL)
    _autosize(al, [22, 18, 30])

    # ----- Sheet Thông số -----
    pa = wb.create_sheet(SHEET_PARAMS)
    pa.append(["station_name", "tag", "param_label", "param_unit",
               "param_low", "param_high", "sort_order", "active"])
    _style_header(pa, 8)
    ex_pa = [
        ["TK-5211A", "052-LI-042B", "Tank level", "mm", 100, 4000, 1, "TRUE"],
        ["PUMP_STATION_6", "052-PG-038", "Áp suất", "bar", 2, 8, 1, "TRUE"],
        ["PUMP_STATION_6", "052-PG-038", "Nhiệt độ", "°C", "", 75, 2, "TRUE"],
        ["PUMP_STATION_6", "052-PG-038", "Chạy/Dừng", "Yes/No", "", "", 3, "TRUE"],
    ]
    _add_rows(pa, ex_pa, _EXAMPLE_FILL)
    _autosize(pa, [18, 16, 18, 12, 12, 12, 11, 10])

    return wb
